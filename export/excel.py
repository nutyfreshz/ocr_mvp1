from __future__ import annotations

from io import BytesIO
from typing import Dict, List

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

RECORD_COLUMNS = [
    "source_file", "page_number", "document_type", "detected_language",
    "citizen_id", "passport_number", "name_native", "surname_native",
    "name_english", "surname_english", "date_of_birth", "sex", "nationality",
    "issue_date", "expiry_date", "issuing_country", "address",
    "overall_confidence", "validation_status", "review_required",
]


def _style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        values = [str(c.value or "") for c in list(column_cells)[:200]]
        width = min(45, max(10, max((len(v) for v in values), default=8) + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def build_excel(records: List[Dict], issues: List[Dict], raw_rows: List[Dict]) -> bytes:
    records_df = pd.DataFrame(records)
    for col in RECORD_COLUMNS:
        if col not in records_df.columns:
            records_df[col] = ""
    records_df = records_df[RECORD_COLUMNS]
    issues_df = pd.DataFrame(issues, columns=["source_file", "page_number", "field", "value", "issue"])
    raw_df = pd.DataFrame(raw_rows, columns=["source_file", "page_number", "model", "text", "confidence", "box"])

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        records_df.to_excel(writer, index=False, sheet_name="RECORDS")
        issues_df.to_excel(writer, index=False, sheet_name="ISSUES")
        raw_df.to_excel(writer, index=False, sheet_name="RAW_OCR")
        for ws in writer.book.worksheets:
            _style_sheet(ws)
        ws = writer.book["RECORDS"]
        status_col = RECORD_COLUMNS.index("validation_status") + 1
        for row in range(2, ws.max_row + 1):
            status = ws.cell(row=row, column=status_col).value
            if status == "REVIEW":
                ws.cell(row=row, column=status_col).fill = PatternFill("solid", fgColor="FFF2CC")
            elif status == "FAIL":
                ws.cell(row=row, column=status_col).fill = PatternFill("solid", fgColor="F4CCCC")
    return output.getvalue()
